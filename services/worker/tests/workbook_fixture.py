"""Real OOXML fixtures written by openpyxl; no customer files or production dependency."""

import base64
import io
import sys

from openpyxl import Workbook


def workbook_bytes(*, incoming=False, formula=False):
    workbook = Workbook()
    notes = workbook.active
    notes.title = "Notes"
    notes.append(["Choose Products, header row 2"])
    sheet = workbook.create_sheet("Products")
    sheet.append(["Supplier October catalogue"])
    sheet.append(["SKU", "Price", "Description", "Currency", "Pack", "UOM"])
    sheet.append(
        [
            "00123",
            "1.002000000000000001" if incoming else "1.001000000000000000",
            "Precision component",
            "GBP",
            1,
            "each",
        ]
    )
    sheet.append([42, 5.25, "Padded identifier", "GBP", 1, "each"])
    sheet["A4"].number_format = "00000"
    sheet.append(["REVIEW", None, "Missing supplier cost", "GBP", 1, "each"])
    if formula:
        sheet["B3"] = "=1+1"
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["Do not import"])
    workbook.create_sheet("Empty")
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


if __name__ == "__main__":
    print(
        base64.b64encode(
            workbook_bytes(incoming="incoming" in sys.argv, formula="formula" in sys.argv)
        ).decode("ascii")
    )
