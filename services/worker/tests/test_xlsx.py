"""Direct XLSX parsing, typed cells, exact decimals and shared reconciliation regressions."""

import io
from dataclasses import asdict
from datetime import date
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook
from workbook_fixture import workbook_bytes

from stockshift_worker.domain.csv_engine import CsvOptions, parse_csv, reconcile
from stockshift_worker.domain.xlsx import WorkbookError, inspect_workbook, parse_xlsx

OPTIONS = {
    **asdict(
        CsvOptions(
            encoding="utf-8-sig",
            delimiter=",",
            decimal_separator=".",
            columns={
                "supplier_sku": "SKU",
                "cost_price": "Price",
                "description": "Description",
                "currency": "Currency",
                "pack_quantity": "Pack",
                "unit": "UOM",
            },
            price_basis="unit",
            tax_basis="net",
        )
    ),
    "format": "xlsx",
    "worksheet": "Products",
    "header_row": 2,
}


def parse(data=None, **settings):
    return parse_xlsx(
        data if data is not None else workbook_bytes(),
        source_file_id=str(uuid4()),
        source_name="catalogue.xlsx",
        options={**OPTIONS, **settings},
    )


def edit(change):
    book = load_workbook(io.BytesIO(workbook_bytes()))
    change(book["Products"])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def patch_part(part, transform):
    output = io.BytesIO()
    with (
        ZipFile(io.BytesIO(workbook_bytes())) as source,
        ZipFile(output, "w", ZIP_DEFLATED) as target,
    ):
        for name in source.namelist():
            data = source.read(name)
            target.writestr(name, transform(data) if name == part else data)
    return output.getvalue()


def test_simple_comparison_and_review_share_engine():
    old, new = parse(), parse(workbook_bytes(incoming=True))
    results = reconcile(old, new)
    assert [r["outcome"] for r in results] == [
        "changed",
        "unchanged",
        "needs_review",
        "needs_review",
    ]
    assert results[0]["cost_delta"] == "0.001000000000000001"


def test_multiple_worksheets_are_discovered_without_guessing():
    metadata = inspect_workbook(workbook_bytes())
    assert [s["name"] for s in metadata["worksheets"]] == ["Notes", "Products", "Hidden", "Empty"]
    assert metadata["worksheets"][2]["state"] == "hidden"
    assert "headers" not in metadata


def test_explicit_header_preview_and_sheet_selection():
    metadata = inspect_workbook(workbook_bytes(), "Products", 2)
    assert metadata["headers"] == ["SKU", "Price", "Description", "Currency", "Pack", "UOM"]
    assert metadata["samples"][0][0] == "00123"
    with pytest.raises(WorkbookError, match="Mapped column"):
        parse(worksheet="Notes", header_row=1)


def test_text_and_simple_zero_mask_identifiers():
    assert [p["supplier_sku"] for p in parse().products][:2] == ["00123", "00042"]


def test_numeric_xml_decimal_precision_is_not_float_coerced():
    data = patch_part(
        "xl/worksheets/sheet2.xml",
        lambda b: b.replace(b"<v>5.25</v>", b"<v>5.250000000000000001</v>"),
    )
    assert parse(data).products[1]["cost_price"] == "5.250000000000000001"
    assert parse(data, decimal_separator=",").products[1]["cost_price"] == "5.250000000000000001"


def test_blank_price_and_other_cells_remain_null():
    product = parse().products[2]
    assert product["cost_price"] is None
    assert any(i["code"] == "missing_value" for i in product["validation_issues"])


def test_invalid_text_price_is_review_not_zero():
    product = parse(edit(lambda s: setattr(s["B3"], "value", "bad price"))).products[0]
    assert product["cost_price"] is None
    assert product["validation_issues"][0]["code"] == "invalid_decimal"


def test_provenance_retains_sheet_physical_row_header_raw_value():
    evidence = parse().evidence[6]
    assert evidence["locator"]["sheet"] == "Products"
    assert evidence["locator"]["row"] == 4
    assert evidence["locator"]["column"] == "SKU"
    assert evidence["raw_text"] == "42"


@pytest.mark.parametrize("worksheet", [None, "Missing", "Hidden", "Empty"])
def test_invalid_sheet_is_actionable(worksheet):
    with pytest.raises(WorkbookError):
        parse(worksheet=worksheet)


@pytest.mark.parametrize("header_row", [None, 0, 201, 1, 1.5, True])
def test_invalid_header_row_rejected(header_row):
    with pytest.raises(WorkbookError):
        parse(header_row=header_row)


def test_duplicate_headers_rejected():
    with pytest.raises(WorkbookError, match="Duplicate header"):
        parse(edit(lambda s: setattr(s["C2"], "value", "SKU")))


def test_missing_mapping_rejected():
    with pytest.raises(WorkbookError, match="Mapped column"):
        parse(columns={"supplier_sku": "Missing", "cost_price": "Price"})


@pytest.mark.parametrize("cached", [False, True])
def test_formulas_are_never_guessed_even_with_cached_value(cached):
    data = workbook_bytes(formula=True)
    if cached:
        output = io.BytesIO()
        with ZipFile(io.BytesIO(data)) as src, ZipFile(output, "w", ZIP_DEFLATED) as dst:
            for name in src.namelist():
                raw = src.read(name)
                if name == "xl/worksheets/sheet2.xml":
                    raw = raw.replace(b"<f>1+1</f><v></v>", b"<f>1+1</f><v>2</v>")
                dst.writestr(name, raw)
        data = output.getvalue()
    with pytest.raises(WorkbookError, match="formula"):
        parse(data)


@pytest.mark.parametrize("value", [date(2026, 10, 5), True, "#DIV/0!"])
def test_dates_booleans_errors_cannot_be_product_prices(value):
    with pytest.raises(WorkbookError):
        parse(edit(lambda s: setattr(s["B3"], "value", value)))


def test_merged_header_layout_rejected():
    with pytest.raises(WorkbookError, match="Merged"):
        parse(edit(lambda s: s.merge_cells("A1:C1")))


@pytest.mark.parametrize(
    "data", [b"broken", b"PK\x03\x04broken", bytes.fromhex("d0cf11e0a1b11ae1") + b"encrypted"]
)
def test_corrupt_and_encrypted_container_actionable(data):
    with pytest.raises(WorkbookError):
        parse(data)


def test_no_worksheets_rejected():
    data = patch_part(
        "xl/workbook.xml",
        lambda b: b[: b.index(b"<sheets>")] + b"<sheets/>" + b[b.index(b"</sheets>") + 9 :],
    )
    with pytest.raises(WorkbookError, match="no worksheets"):
        parse(data)


@pytest.mark.parametrize("value", [1234567890123456, 1.5, 1e20])
def test_ambiguous_numeric_identifier_rejected(value):
    with pytest.raises(WorkbookError, match="ambiguous numeric"):
        parse(edit(lambda s: setattr(s["A3"], "value", value)))


def test_xml_entity_attack_rejected():
    data = patch_part(
        "xl/workbook.xml", lambda _: b'<!DOCTYPE foo [<!ENTITY x "boom">]><foo>&x;</foo>'
    )
    with pytest.raises(WorkbookError, match="Corrupt"):
        parse(data)


def test_mixed_csv_xlsx_uses_identical_record_contract():
    syntax = CsvOptions(
        encoding="utf-8-sig",
        delimiter=",",
        decimal_separator=".",
        columns={"supplier_sku": "SKU", "cost_price": "Price", "description": "Description"},
        currency="GBP",
        unit="each",
        pack_quantity="1",
        price_basis="unit",
        tax_basis="net",
    )
    csv = parse_csv(
        b"SKU,Price,Description\n00123,1.001000000000000000,Precision component\n",
        source_file_id=str(uuid4()),
        source_name="old.csv",
        options=syntax,
    )
    results = reconcile(csv, parse(workbook_bytes(incoming=True)))
    assert results[0]["cost_delta"] == "0.001000000000000001"


def test_blank_date_styled_cell_is_null_not_a_date():
    def change(sheet):
        sheet["B3"].value = None
        sheet["B3"].number_format = "yyyy-mm-dd"

    assert parse(edit(change)).products[0]["cost_price"] is None


def test_sheet_identity_is_part_of_deterministic_record_ids():
    file_id = str(uuid4())
    data = workbook_bytes()
    a = parse_xlsx(data, source_file_id=file_id, source_name="a.xlsx", options=OPTIONS)
    b = parse_xlsx(data, source_file_id=file_id, source_name="a.xlsx", options=OPTIONS)
    assert a == b
    assert len(set(p["record_id"] for p in a.products)) == 3


def test_no_visible_worksheets_is_actionable():
    data = patch_part("xl/workbook.xml", lambda b: b.replace(b'state="visible"', b'state="hidden"'))
    with pytest.raises(WorkbookError, match="no visible worksheets"):
        parse(data)


def test_invalid_shared_string_reference_rejected():
    data = patch_part(
        "xl/worksheets/sheet2.xml",
        lambda b: b.replace(
            b'<c r="A3" t="inlineStr"><is><t>00123</t></is></c>', b'<c r="A3" t="s"><v>-1</v></c>'
        ),
    )
    with pytest.raises(WorkbookError, match="shared text reference"):
        parse(data)
